from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import pytest

from oslm.ingest.crawler import Crawler, Source, read_sources
from oslm.ingest.loaders import html_to_text, load_local, load_pdf, load_pptx
from oslm.ingest.manifest import Document, Manifest, ManifestEntry, Segment, content_hash, make_doc_id
from oslm.ingest.run import ingest_course

ROOT = Path(__file__).resolve().parent.parent


@dataclass
class FakeResponse:
    url: str
    status_code: int = 200
    content: bytes = b""
    headers: dict[str, str] = field(default_factory=lambda: {"Content-Type": "text/html"})

    @property
    def text(self) -> str:
        return self.content.decode("utf-8")


class FakeSession:
    def __init__(self, pages: dict[str, bytes | int]) -> None:
        self.pages = pages
        self.calls: list[str] = []
        self.headers: dict[str, str] = {}

    def get(self, url: str, timeout: float) -> FakeResponse:
        self.calls.append(url)
        page = self.pages.get(url, 404)
        if isinstance(page, int):
            return FakeResponse(url, status_code=page)
        return FakeResponse(url, content=page)


def make_crawler(tmp_path: Path, pages: dict[str, bytes | int], domains: set[str]) -> tuple[Crawler, FakeSession]:
    session = FakeSession(pages)
    crawler = Crawler(
        cache_dir=tmp_path / "cache",
        user_agent="OSlm-test",
        allowed_domains=domains,
        delay_s=0,
        retries=1,
        session=session,  # type: ignore[arg-type]
        sleep=lambda s: None,
    )
    return crawler, session


def source(url: str, depth: int = 0, prefix: str = "/wiki/") -> Source:
    return Source(url, "Test", "html", "CC", 2, depth, prefix, ("t",))


class TestCrawler:
    def test_robots_disallow_is_respected(self, tmp_path: Path) -> None:
        pages = {
            "https://ex.org/robots.txt": b"User-agent: *\nDisallow: /private/\n",
            "https://ex.org/private/a": b"<html><body>secret</body></html>",
        }
        crawler, session = make_crawler(tmp_path, pages, {"ex.org"})
        assert crawler.fetch("https://ex.org/private/a") is None
        assert "https://ex.org/private/a" not in session.calls

    def test_off_prefix_and_off_domain_links_skipped(self, tmp_path: Path) -> None:
        start = "https://ex.org/wiki/Start"
        html = (
            b"<html><body><main><p>Intro</p>"
            b'<a href="/wiki/Inside">in</a><a href="/other/Outside">out</a>'
            b'<a href="https://evil.com/wiki/X">ext</a><a href="#frag">frag</a></main></body></html>'
        )
        pages = {
            "https://ex.org/robots.txt": b"User-agent: *\nAllow: /\n",
            start: html,
            "https://ex.org/wiki/Inside": b"<html><body><p>inside page</p></body></html>",
        }
        crawler, session = make_crawler(tmp_path, pages, {"ex.org"})
        results = crawler.crawl(source(start, depth=1))
        assert [r.url for r in results] == [start, "https://ex.org/wiki/Inside"]
        assert not any("Outside" in c or "evil" in c for c in session.calls)

    def test_max_depth_zero_does_not_follow_links(self, tmp_path: Path) -> None:
        start = "https://ex.org/wiki/Start"
        pages = {
            "https://ex.org/robots.txt": b"",
            start: b'<html><body><a href="/wiki/Inside">in</a></body></html>',
            "https://ex.org/wiki/Inside": b"<html></html>",
        }
        crawler, session = make_crawler(tmp_path, pages, {"ex.org"})
        assert len(crawler.crawl(source(start, depth=0))) == 1
        assert "https://ex.org/wiki/Inside" not in session.calls

    def test_cache_prevents_refetch(self, tmp_path: Path) -> None:
        url = "https://ex.org/wiki/A"
        pages = {"https://ex.org/robots.txt": b"", url: b"<html><body>a</body></html>"}
        crawler, session = make_crawler(tmp_path, pages, {"ex.org"})
        first = crawler.fetch(url)
        n_calls = len(session.calls)
        crawler2, session2 = make_crawler(tmp_path, pages, {"ex.org"})
        second = crawler2.fetch(url)
        assert first and second and second.from_cache and second.body == first.body
        assert session2.calls == [] and n_calls == 2

    def test_retries_then_gives_up(self, tmp_path: Path) -> None:
        url = "https://ex.org/wiki/Flaky"
        crawler, session = make_crawler(tmp_path, {"https://ex.org/robots.txt": b"", url: 503}, {"ex.org"})
        with pytest.raises(RuntimeError):
            crawler.fetch(url)
        assert session.calls.count(url) == 2

    def test_unlisted_domain_never_requested(self, tmp_path: Path) -> None:
        crawler, session = make_crawler(tmp_path, {}, {"ex.org"})
        assert crawler.fetch("https://other.org/x") is None
        assert session.calls == []

    def test_crawl_delay_from_robots(self, tmp_path: Path) -> None:
        sleeps: list[float] = []
        pages = {
            "https://ex.org/robots.txt": b"User-agent: *\nCrawl-delay: 2\n",
            "https://ex.org/wiki/A": b"<html></html>",
            "https://ex.org/wiki/B": b"<html></html>",
        }
        session = FakeSession(pages)
        clock = iter(range(0, 1000)).__next__
        crawler = Crawler(
            tmp_path / "c",
            "ua",
            {"ex.org"},
            delay_s=1,
            retries=0,
            session=session,  # type: ignore[arg-type]
            sleep=sleeps.append,
            clock=lambda: float(clock()),
        )
        crawler.fetch("https://ex.org/wiki/A")
        crawler.fetch("https://ex.org/wiki/B")
        assert sleeps and max(sleeps) > 0.5

    def test_sources_csv_is_valid(self) -> None:
        sources = read_sources(ROOT / "sources.csv")
        assert len(sources) >= 30
        assert len({s.url for s in sources}) == len(sources)
        assert all(s.type in ("html", "pdf") and 1 <= s.trust_tier <= 3 for s in sources)
        assert all(s.url.startswith("https://") and s.path_prefix.startswith("/") for s in sources)
        assert not any("man7.org" in s.url for s in sources)


class TestHtml:
    def test_strips_navigation_and_keeps_code(self) -> None:
        html = """<html><head><title>Spinlock - Wikipedia</title></head><body>
        <nav>Main menu Donate</nav><div id="mw-content-text"><h1>Spinlock</h1>
        <p>A spinlock waits in a loop.</p><div class="navbox">Navbox junk</div>
        <pre>while (test_and_set(&amp;lock)) ;</pre>
        <h2>See also</h2><p>Unrelated list</p></div><footer>Privacy policy</footer></body></html>"""
        title, text = html_to_text(html)
        assert title == "Spinlock"
        assert "spinlock waits" in text
        assert "```\nwhile (test_and_set(&lock)) ;\n```" in text
        for junk in ("Main menu", "Navbox junk", "Privacy policy", "Unrelated list"):
            assert junk not in text


def _tiny_pdf() -> bytes:
    import fitz

    doc = fitz.open()
    for i in range(2):
        page = doc.new_page()
        page.insert_text((72, 72), f"Semaphore page {i + 1}: wait and signal operations.")
    data = doc.tobytes()
    doc.close()
    return data


class TestLoaders:
    def test_pdf_loader_extracts_pages(self) -> None:
        _, segments = load_pdf(_tiny_pdf())
        assert [s.loc for s in segments] == ["p.1", "p.2"]
        assert "wait and signal" in segments[1].text

    def test_pptx_loader_extracts_text_and_notes(self, tmp_path: Path) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Critical Section"
        slide.placeholders[1].text = "Only one process may execute it at a time"
        slide.notes_slide.notes_text_frame.text = "Mention entry and exit sections"
        rows = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
        rows.cell(0, 0).text, rows.cell(0, 1).text = "Op", "Effect"
        rows.cell(1, 0).text, rows.cell(1, 1).text = "P", "decrement"
        path = tmp_path / "lecture.pptx"
        prs.save(str(path))

        _, segments = load_pptx(path)
        assert segments[0].loc == "slide 1"
        text = segments[0].text
        assert "## Critical Section" in text and "one process" in text
        assert "Speaker notes: Mention entry" in text and "| P | decrement |" in text

    def test_text_loader(self, tmp_path: Path) -> None:
        path = tmp_path / "notes.md"
        path.write_text("# Deadlock notes\n\nFour conditions: mutual exclusion…", encoding="utf-8")
        doc_type, title, segments = load_local(path)
        assert doc_type == "md" and title == "Deadlock notes" and "…" in segments[0].text


def _entry(doc_id: str, url: str, segments: list[Segment]) -> ManifestEntry:
    return ManifestEntry(doc_id, "S", url, "html", "CC", 2, "T", "now", content_hash(segments))


class TestManifest:
    def test_no_duplicate_hashes(self, tmp_path: Path) -> None:
        m = Manifest(tmp_path)
        segs = [Segment("same text")]
        assert m.add(Document(_entry("a", "https://x/a", segs), segs))
        assert not m.add(Document(_entry("b", "https://x/b", segs), segs))
        assert not m.add(Document(_entry("a", "https://x/a", segs), segs))
        hashes = [e.sha256 for e in Manifest(tmp_path).entries.values()]
        assert len(hashes) == len(set(hashes)) == 1

    def test_changed_content_replaces_entry(self, tmp_path: Path) -> None:
        m = Manifest(tmp_path)
        s1, s2 = [Segment("v1")], [Segment("v2")]
        m.add(Document(_entry("a", "https://x/a", s1), s1))
        assert m.add(Document(_entry("a2", "https://x/a", s2), s2))
        reloaded = Manifest(tmp_path)
        assert list(reloaded.entries) == ["a2"] and reloaded.load("a2").segments[0].text == "v2"

    def test_doc_id_is_readable_and_stable(self) -> None:
        a = make_doc_id("OSTEP", "Locks", "https://x/locks.pdf")
        assert a == make_doc_id("OSTEP", "Locks", "https://x/locks.pdf")
        assert a.startswith("ostep-locks-")

    def test_course_ingest_is_idempotent(self, tmp_path: Path) -> None:
        course = tmp_path / "course"
        course.mkdir()
        (course / "notes.txt").write_text("Semaphores: P and V.", encoding="utf-8")
        (course / "slides.pdf").write_bytes(_tiny_pdf())
        (course / "ignore.xyz").write_text("nope", encoding="utf-8")
        m = Manifest(tmp_path / "raw")
        assert ingest_course(m, course) == (2, 0)
        assert ingest_course(Manifest(tmp_path / "raw"), course) == (0, 2)
