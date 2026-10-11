from __future__ import annotations

import io
import logging
import re
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag

from oslm.ingest.manifest import Segment

log = logging.getLogger(__name__)

_STRIP_TAGS = ["script", "style", "nav", "footer", "aside", "header", "form", "noscript", "svg", "button"]
_STRIP_SELECTORS = [
    ".navbox",
    ".vertical-navbox",
    ".sidebar",
    ".mw-editsection",
    ".reference",
    ".reflist",
    ".references",
    "#toc",
    ".toc",
    ".noprint",
    ".ambox",
    ".hatnote",
    ".mw-jump-link",
    ".catlinks",
    "#siteSub",
    "#contentSub",
    ".printfooter",
    ".mw-indicators",
    ".metadata",
    ".shortdescription",
    ".infobox",
    "#mw-navigation",
    ".man-page-footer",
    ".nav",
    "#footer",
    "table.head",
    "table.foot",
]
_CONTENT_SELECTORS = ["#mw-content-text", "main", "article", "#content", ".man-page", "body"]
_TAIL_SECTIONS = re.compile(
    r"^(see also|references|external links|further reading|bibliography|colophon|copyright)$", re.I
)


def _pre_to_fence(pre: Tag) -> str:
    code = pre.get_text()
    return f"\n```\n{code.rstrip()}\n```\n"


def html_to_text(html: str | bytes) -> tuple[str, str]:
    """Returns (title, text) with nav/footer/sidebars removed, headings as `#` lines and <pre> as fenced code."""
    soup = BeautifulSoup(html, "lxml")
    h1 = soup.find("h1")
    title = h1.get_text(" ", strip=True) if h1 else ""
    # Man pages start with a section heading such as NAME; their <title> names the page.
    if (not title or title.isupper()) and soup.title:
        title = soup.title.get_text(" ", strip=True)
    title = re.sub(r"\s*[-–—|]\s*(Wikipedia|Arch manual pages).*$", "", title)

    root = next((soup.select_one(sel) for sel in _CONTENT_SELECTORS if soup.select_one(sel)), soup)
    for tag in root.find_all(_STRIP_TAGS):
        tag.decompose()
    for sel in _STRIP_SELECTORS:
        for tag in root.select(sel):
            tag.decompose()

    for pre in root.find_all("pre"):
        pre.replace_with(NavigableString(_pre_to_fence(pre)))
    for level in range(1, 5):
        for h in root.find_all(f"h{level}"):
            h.replace_with(NavigableString(f"\n\n{'#' * level} {h.get_text(' ', strip=True)}\n\n"))
    for li in root.find_all("li"):
        li.insert_before(NavigableString("\n- "))
    for block in root.find_all(["p", "div", "dd", "dt", "tr", "br", "table", "ul", "ol", "dl"]):
        block.insert_after(NavigableString("\n"))
    for cell in root.find_all(["td", "th"]):
        cell.insert_after(NavigableString(" | "))

    text = root.get_text()
    text = _drop_tail_sections(text)
    return title, text


def _drop_tail_sections(text: str) -> str:
    out: list[str] = []
    skipping = False
    skip_level = 0
    for line in text.splitlines():
        m = re.match(r"^(#{1,4}) (.+)$", line.strip())
        if m:
            level = len(m.group(1))
            if _TAIL_SECTIONS.match(m.group(2).strip()):
                skipping, skip_level = True, level
                continue
            if skipping and level <= skip_level:
                skipping = False
        if not skipping:
            out.append(line)
    return "\n".join(out)


_JUNK_PDF_TITLE = re.compile(r"(\.(dvi|tex|pdf|docx?)$|^untitled|^microsoft word|^slide ?\d*$)", re.I)


def pdf_title(metadata_title: str, first_page: str) -> str:
    """Metadata title unless it is a build artefact like "single.dvi"; then the first real line of page 1."""
    title = metadata_title.strip()
    if title and not _JUNK_PDF_TITLE.search(title):
        if " " not in title and re.search(r"[a-z][A-Z]", title):
            title = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", title)
        return title
    for line in first_page.splitlines():
        line = line.strip()
        if len(re.findall(r"[A-Za-z]", line)) >= 3:
            return line[:120]
    return ""


def load_pdf(data: bytes) -> tuple[str, list[Segment]]:
    import fitz  # PyMuPDF

    segments: list[Segment] = []
    title = ""
    try:
        with fitz.open(stream=data, filetype="pdf") as doc:
            title = (doc.metadata or {}).get("title", "") or ""
            for i, page in enumerate(doc, start=1):
                text = page.get_text("text")
                has_tables = False
                try:
                    has_tables = bool(page.find_tables().tables)
                except Exception:
                    log.debug("find_tables failed on page %d", i)
                if has_tables or len(text.strip()) < 20:
                    text = _pdfplumber_page(data, i - 1, text, include_tables=has_tables)
                segments.append(Segment(text=text, loc=f"p.{i}"))
    except Exception as exc:
        log.warning("PyMuPDF failed (%s); falling back to pdfplumber", exc)
        segments = _pdfplumber_all(data)
    return pdf_title(title, segments[0].text if segments else ""), segments


def _table_to_markdown(rows: list[list[str | None]]) -> str:
    clean = [[(c or "").replace("\n", " ").strip() for c in row] for row in rows if row]
    if not clean:
        return ""
    width = max(len(r) for r in clean)
    clean = [r + [""] * (width - len(r)) for r in clean]
    lines = ["| " + " | ".join(clean[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in clean[1:]]
    return "\n".join(lines)


def _pdfplumber_page(data: bytes, index: int, fallback: str, include_tables: bool) -> str:
    import pdfplumber

    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            page = pdf.pages[index]
            text = page.extract_text() or fallback
            if include_tables:
                tables = [_table_to_markdown(t) for t in page.extract_tables()]
                text += "\n\n" + "\n\n".join(t for t in tables if t)
            return text
    except Exception as exc:
        log.debug("pdfplumber failed on page %d: %s", index + 1, exc)
        return fallback


def _pdfplumber_all(data: bytes) -> list[Segment]:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return [Segment(text=p.extract_text() or "", loc=f"p.{i}") for i, p in enumerate(pdf.pages, start=1)]


def load_pptx(path: Path) -> tuple[str, list[Segment]]:
    from pptx import Presentation

    prs = Presentation(str(path))
    segments: list[Segment] = []
    title = prs.core_properties.title or ""
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            heading = slide.shapes.title.text_frame.text.strip()
            if heading:
                parts.append(f"## {heading}")
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                parts.append(_table_to_markdown(rows))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                parts.append(f"Speaker notes: {notes}")
        if parts:
            segments.append(Segment(text="\n\n".join(parts), loc=f"slide {i}"))
    return title, segments


def load_text(path: Path) -> tuple[str, list[Segment]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    first = next((ln.lstrip("# ").strip() for ln in text.splitlines() if ln.strip()), path.stem)
    return first[:120], [Segment(text=text)]


LOCAL_SUFFIXES = {".pdf", ".pptx", ".txt", ".md"}


def load_local(path: Path) -> tuple[str, str, list[Segment]]:
    """Returns (type, title, segments) for a file dropped into data/course/."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        title, segs = load_pdf(path.read_bytes())
        return "pdf", title or path.stem, segs
    if suffix == ".pptx":
        title, segs = load_pptx(path)
        return "pptx", title or path.stem, segs
    if suffix in (".txt", ".md"):
        title, segs = load_text(path)
        return suffix.lstrip("."), title, segs
    raise ValueError(f"unsupported file type: {path.name}")
